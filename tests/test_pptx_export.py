import io

from pptx import Presentation

from app.tools.pptx_export import _parse, generate_pptx

SAMPLE = """TITLE: CRM readiness in mid-market pharma
SUBTITLE: TPDL point of view
SLIDE 1: The shift / New commercial leaders inherit fragmented CRM
SLIDE 2: The cost
- Duplicated data across markets
- No single customer view
SECTION 3 — What good looks like
- One operating model
- Activation-ready data
"""


def test_parse_cover_and_slides():
    title, subtitle, slides = _parse(SAMPLE, "fallback-subject", "label")
    assert title == "CRM readiness in mid-market pharma"
    assert subtitle == "TPDL point of view"
    assert len(slides) == 3
    assert slides[0]["title"] == "The shift"
    assert "inherit" in slides[0]["body"][0]
    assert slides[1]["body"] == ["Duplicated data across markets", "No single customer view"]
    assert slides[2]["title"] == "What good looks like"


def test_generate_pptx_valid_and_slide_count():
    data = generate_pptx("CRM readiness", SAMPLE, "TPDL Deck")
    assert isinstance(data, bytes) and len(data) > 2000
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 4  # 1 cover + 3 content
    joined = "\n".join(
        sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame
    )
    assert "CRM readiness in mid-market pharma" in joined
    assert "The shift" in joined
    assert "Duplicated data across markets" in joined


def test_generate_pptx_fallback_without_markers():
    data = generate_pptx("Plain topic", "just some prose\nmore prose", "TPDL Deck")
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 2  # cover + one fallback content slide
