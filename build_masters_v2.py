#!/usr/bin/env python3
"""Masters Branding / Stratégie de marque / Partnerships — Espagne & online — ≤15k€.
Présentation calquée sur l'ancien Excel Masters_Brand_Luxe_Europe_2026.xlsx."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NAVY = "1F3864"
GREEN_DARK = "375623"
GREEN_MID = "A9D08E"
GREEN_LIGHT = "C6EFCE"
YELLOW = "FFF2CC"
GREY = "D9D9D9"
RED_LIGHT = "F8CBAD"
WHITE = "FFFFFF"
LINK_BLUE = "0563C1"

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(c):
    c.font = Font(name=FONT, bold=True, color=WHITE, size=10)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border

def style_cell(c, fill=None, bold=False, size=9, color="000000", align="left"):
    c.font = Font(name=FONT, bold=bold, size=size, color=color)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=True)
    c.border = border

wb = Workbook()

# ============================================================
# ONGLET 1 — COMPARATIF (présentation ancienne : en-têtes ligne 1)
# ============================================================
ws = wb.active
ws.title = "Comparatif"

headers = [
    ("Master", 34),
    ("École", 26),
    ("Pays / lieu", 18),
    ("Réputation de l'école", 28),
    ("Classements & reviews", 26),
    ("Durée", 12),
    ("Modalité (détail)", 30),
    ("OK avec full-time job", 18),
    ("Langue(s)", 12),
    ("Prix", 24),
    ("Critères d'admission", 24),
    ("Deadline / période candidature", 20),
    ("Stage", 22),
    ("Alumni / employeurs", 26),
    ("Particularité / à savoir", 28),
    ("Lien direct", 22),
    ("Verdict", 26),
]
for j, (name, w) in enumerate(headers, start=1):
    style_header(ws.cell(1, j, name))
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 60

# Each program: 16 data fields (cols 1-16) + verdict text + verdict fill
programs = [
    {
        "fill": GREEN_MID,
        "verdict": "★★★★★ RECO N°1 — branding pur, prix ferme 12k, remote possible depuis Madrid",
        "d": [
            "Master in Brand Strategy & Creative Brand Management",
            "UPF-BSM — Barcelona School of Management (Universitat Pompeu Fabra)",
            "Espagne / Barcelone (+ option remote synchrone)",
            "TRÈS reconnue — UPF = top 1-3 universités d'Espagne (QS). BSM = business school officielle de l'UPF, très cotée en marketing/branding.",
            "UPF ~top 3 ES (QS). Diplôme universitaire de poids. Réseau brand & agences de Barcelone.",
            "10 mois (60 ECTS)",
            "Présentiel Vendredi + Samedi (hybride). Option À DISTANCE SYNCHRONE — max 15 places/promo, même horaire live → suivable depuis Madrid.",
            "OUI — pensé pour actifs (ven. soir + samedi)",
            "Espagnol",
            "12 000 € (vérifié 3-0)",
            "Dossier + CV + entretien. Profil marketing/com/business ou créa.",
            "Rentrée automne — postuler dès le printemps (places remote limitées à 15 → tôt).",
            "Proposé (réseau d'entreprises BSM, agences & marques).",
            "Brand managers, brand directors, strategic planners — départements marketing, agences de pub, cabinets de conseil en marque.",
            "LE master le plus centré stratégie de marque d'Espagne + format pour actifs. Réseau brand/agences de Barcelone.",
            "https://www.bsm.upf.edu/en/master-in-brand-strategy-and-creative-brand-management",
        ],
    },
    {
        "fill": GREEN_MID,
        "verdict": "★★★★ RECO N°2 — luxe pur ; online 11 450 € pile dans le budget",
        "d": [
            "MBA en Marcas de Lujo (MBA marques de luxe)",
            "ELLE Education (Madrid)",
            "Espagne / Madrid (+ online)",
            "Marque média ELLE (mode/luxe) — bonne notoriété sectorielle. École spécialisée, pas université classée QS : la valeur = réseau & branding secteur.",
            "Pas de classement académique ; réputation sectorielle mode/luxe. Vérifier les avis d'anciens.",
            "≈ 1 an",
            "Semi-présentiel Madrid : 28 week-ends — Ven 17h00-20h30 + Sam 09h30-14h00. OU 100% online flexible 24/7.",
            "OUI — weekend ou online",
            "Espagnol",
            "Semi-présentiel 14 950 € · Executive 14 450 € · Online 11 450 € (inclut matrícula, supports, cours) — vérifié",
            "Dossier + entretien. Profil mode/luxe/marketing apprécié.",
            "Rentrée annuelle — contacter admissions pour les dates 2026/27.",
            "Proposé (réseau secteur mode/luxe via ELLE).",
            "Postes brand & communication chez marques de luxe et médias mode.",
            "Branding via la marque média ELLE ; format weekend Madrid taillé pour les actifs du luxe.",
            "https://elle.education/programas/mba/marcas-lujo/",
        ],
    },
    {
        "fill": GREEN_MID,
        "verdict": "★★★★ RECO N°3 — master OFFICIEL mode online (demander le prix)",
        "d": [
            "Máster Universitario en Marketing y Comunicación de Moda (Online)",
            "Universidad Europea de Madrid / Creative Campus",
            "Espagne / Madrid (+ online)",
            "Reconnue — grande université privée de Madrid, diplôme universitaire OFFICIEL (60 ECTS). Réseau mode via Creative Campus.",
            "Université privée correcte ; diplôme OFFICIEL (vs titre propre). Bonne visibilité dans la mode.",
            "10-12 mois (60 ECTS) — vérifié",
            "100% online avec classes live, OU « Online & On Campus » = seulement 4 samedis sur le campus Madrid — vérifié.",
            "OUI — online + max 4 samedis",
            "Espagnol (vérifié)",
            "⚠ Non publié (gated). Fourchette UE ≈ 6 000-20 000 €. Promo -23% jusqu'au 15 juin. DEMANDER le prix exact.",
            "Bac+3 + test de compétences + entretien.",
            "Rentrée octobre 2026 (parfois avril). Admission au fil de l'eau.",
            "OBLIGATOIRE — 6 ECTS prácticas externes (inclut shootings Harper's Bazaar, visites) — vérifié.",
            "Brand Manager, Marketing Director (CMO), Communication Director — secteur mode.",
            "Master OFFICIEL spécialisé mode, quasi 100% online + stage obligatoire. Bon rapport reconnaissance/flexibilité.",
            "https://creativecampus.universidadeuropea.com/master-marketing-moda-online/",
        ],
    },
    {
        "fill": GREEN_LIGHT,
        "verdict": "★★★ Très bon — officiel + 100% online + abordable",
        "d": [
            "Máster Universitario en Dirección y Gestión de Marca / Branding",
            "UNIR — Universidad Internacional de La Rioja",
            "Espagne / 100% online",
            "Normale-à-bonne — plus grande université online d'Espagne, diplôme OFFICIEL (ANECA). Réputation académique moyenne mais diplôme valable.",
            "Pas de top classement ; diplôme officiel reconnu. Seul master online OFFICIEL de branding en Espagne.",
            "1 an / 60 ECTS — vérifié",
            "100% online interactif. Examens online et/ou présentiel — vérifié.",
            "OUI — 100% online, conçu pour actifs",
            "Espagnol",
            "⚠ Non publié. Estimé ≈ 4 000-10 000 €. Promo -30% (~2 juin). Comptant ou mensualités. DEMANDER.",
            "Diplôme lié (éco/marketing/com/pub) OU 2 ans d'expérience pro pertinente.",
            "Rentrée novembre 2026.",
            "Inclus — 6 ECTS prácticas externes + mémoire (TFM) 12 ECTS — vérifié.",
            "Postes branding / brand management ; profils en poste ou en reconversion.",
            "Seul master OFFICIEL 100% online de branding en Espagne. Flexibilité max, prix contenu. Marque académique plus modeste.",
            "https://www.unir.net/marketing-comunicacion/master-branding/",
        ],
    },
    {
        "fill": GREEN_LIGHT,
        "verdict": "★★★ Bon — diplôme court le moins cher (ce n'est PAS un master)",
        "d": [
            "Diploma en Marketing y Comunicación de Moda y Lujo",
            "ELLE Education (Madrid)",
            "Espagne / Madrid (+ online)",
            "Marque ELLE (mode/luxe). Format diplôme (pas master) → moins de poids académique mais bon pour le réseau & le secteur.",
            "Réputation sectorielle ; format diplôme court.",
            "Plus court qu'un MBA (diplôme)",
            "Semipresencial : online + sessions weekend Madrid (ven. après-midi / sam. matin). OU 100% online flexible 24/7.",
            "OUI — weekend ou online",
            "Espagnol",
            "Semi-présentiel 4 499 € · Executive 4 199 € · Online 3 119 € (+ 600 € inscription) — vérifié",
            "Dossier. Accessible.",
            "Rentrée annuelle — contacter admissions.",
            "Proposé (réseau ELLE).",
            "Marketing/com mode & luxe ; tremplin ou complément.",
            "Le moins cher pour entrer dans le branding mode/luxe à Madrid. Option « test » avant un master.",
            "https://elle.education/programas/diplomas/moda-y-lujo/",
        ],
    },
    {
        "fill": GREEN_LIGHT,
        "verdict": "★★★ Bon — marque UCM + pas cher, mais digital > branding",
        "d": [
            "Máster en Marketing Digital",
            "Universidad Complutense de Madrid (UCM)",
            "Espagne / Madrid (+ online)",
            "TRÈS reconnue — grande université publique de Madrid, classée QS. Mais ce master est marketing DIGITAL, pas branding pur.",
            "UCM classée QS (université publique de référence). Marque forte.",
            "≈ 1 an",
            "Présentiel weekend (ven. après-midi + sam. matin) OU 100% online flexible — vérifié.",
            "OUI — weekend ou online",
            "Espagnol",
            "Présentiel 6 850 € · Semi-présentiel 5 725 € · Online 4 755 € (+ 40 € pré-inscription) 2026/27 — vérifié",
            "Bac+3. Dossier.",
            "Rentrée annuelle (automne).",
            "Variable.",
            "Profils marketing digital ; marque UCM solide.",
            "Marque universitaire forte + prix très bas. Mais marketing digital ≠ branding/partnerships pur.",
            "https://www.mastermarketingdigitalucm.com/",
        ],
    },
    {
        "fill": YELLOW,
        "verdict": "⚠ Sous 15k MAIS format soir-semaine difficile avec un job",
        "d": [
            "Master in Fashion Companies Management / Dirección de Empresas de Moda",
            "IED Madrid — Istituto Europeo di Design",
            "Espagne / Madrid (Gran Vía)",
            "IED = école de design/mode reconnue internationalement. Bonne marque mode (réseau créatif).",
            "Marque mode/design forte à l'international.",
            "9 mois (60 ECTS) — vérifié",
            "PRÉSENTIEL COURS DU SOIR (« tardes ») à Madrid. ≈ 18h30-22h30, jusqu'à 5 soirs/sem (non confirmé). PAS weekend, PAS online.",
            "⚠ DIFFICILE — soirs en semaine jusqu'à 5j = lourd. Vérifier le nb réel de soirs.",
            "Espagnol",
            "13 200 € — VÉRIFIÉ (page officielle IED, sous 15k). Frais d'inscription éventuels en sus.",
            "Diplômés droit/gestion OU pros mode/marketing/com. CV + lettre + portfolio + entretien.",
            "Rentrée octobre 2026.",
            "Pas de stage formel annoncé (projets réels, visites, projet final) — à confirmer.",
            "Postes management dans des maisons de mode.",
            "Forte marque mode/design + prix sous 15k. Seul frein réel = format soir en semaine.",
            "https://www.ied.es/cursos/madrid/master/direccion-de-empresas-de-moda",
        ],
    },
]

r = 2
for p in programs:
    for j, val in enumerate(p["d"], start=1):
        c = ws.cell(r, j)
        if j == 16:  # lien cliquable
            c.value = "Ouvrir la page ↗"
            c.hyperlink = val
            c.font = Font(name=FONT, size=9, color=LINK_BLUE, underline="single")
            c.fill = PatternFill("solid", fgColor=WHITE)
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            c.border = border
        else:
            c.value = val
            fill = None
            v = str(val)
            if j == 10 and ("⚠" in v or "non publié" in v.lower()):
                fill = YELLOW
            style_cell(c, fill=fill, size=9)
    vc = ws.cell(r, 17, p["verdict"])
    style_cell(vc, fill=p["fill"], bold=True, size=9)
    ws.row_dimensions[r].height = 150
    r += 1

ws.freeze_panes = "C2"

# Légende sous le tableau
leg = ws.cell(r + 1, 1, "Légende — Vert foncé/moyen = recommandé · Vert clair = bon · Jaune = prix non publié ou point d'attention.  « vérifié » = fait confirmé par la recherche (vote 3-0).  Clique sur « Ouvrir la page ↗ » pour le lien direct.")
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=len(headers))
leg.font = Font(name=FONT, bold=True, size=9, italic=True)
leg.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[r + 1].height = 30

# ============================================================
# ONGLET 2 — ÉCARTÉS & HORS BUDGET
# ============================================================
ws2 = wb.create_sheet("Écartés & hors budget")
heads2 = [("Programme", 36), ("École", 26), ("Prix", 20), ("Raison de l'écart", 58), ("Lien", 18)]
for j, (h, w) in enumerate(heads2, start=1):
    style_header(ws2.cell(1, j, h))
    ws2.column_dimensions[get_column_letter(j)].width = w
ws2.row_dimensions[1].height = 30

ecartes = [
    ("Máster Internacional en Marketing de Moda y Lujo (MML)", "ESIC (Madrid & Barcelone)", "20 400 € (vérifié)",
     "TRÈS aligné (mode & luxe, weekend ven+sam, semaine à Marangoni Paris incluse, ESIC bien classé en marketing) MAIS prix VÉRIFIÉ = 20 400 € → dépasse les 15 000 €. Título propio (non officiel). Jusqu'à 30% d'aide possible → net ~14k à négocier si vraiment intéressée.",
     "https://www.esic.edu/master-y-postgrado/master-internacional-en-marketing-de-moda-y-lujo-mml"),
    ("Executive Master in Strategic Marketing & Communication", "IE Business School (Madrid)", "≈ 35 000 €+ (estimé)",
     "ÉCOLE D'ÉLITE (top 3-5 Europe, Madrid & online, 15 mois, anglais) MAIS prix probable > 35 000 € → bien au-dessus du plafond de 15 000 €. Option « rêve » seulement si l'entreprise/financement paie.",
     "https://www.ie.edu/business-school/programs/masters/executive-master-in-strategic-marketing-communication/"),
    ("Fashion & Luxury Brand Management", "Istituto Marangoni (Milan/Florence/Paris/Londres/Dubaï)", "30 600-31 000 € + 5 500 €",
     "Full-time PRÉSENTIEL uniquement, PAS en Espagne, PAS online/weekend → incompatible job full-time à Madrid. Prix très au-dessus du budget. (Vérifié 3-0)",
     "https://www.istitutomarangoni.com/en/fashion-courses/master/fashion-business/fashion-luxury-brand-management"),
    ("MA International Luxury Business (Online)", "Vogue College of Fashion", "≈ 34 950 £ (≈ 40 000 €+)",
     "Online et pensé pour actifs du luxe MAIS prix ≈ 40 000 €+ → très au-dessus du budget.",
     "https://www.voguecollege.com/courses/ma-in-luxury-business-online/"),
    ("Divers programmes", "EU Business School / GBSB Global (Madrid/Barcelone)", "Variable",
     "Très visibles en pub MAIS diplômes peu valorisés selon avis d'anciens. À éviter sauf raison précise — marque académique faible.",
     "—"),
    ("Masters full-time présentiel semaine", "IE / ESADE / IESE (formats full-time)", "Élevé",
     "Excellentes écoles mais formats full-time en semaine = incompatibles avec un job full-time. Seuls online/weekend/executive sont retenus.",
     "—"),
]
r = 2
for prog, ecole, prix, raison, lien in ecartes:
    style_cell(ws2.cell(r, 1, prog), fill=RED_LIGHT, bold=True, size=9)
    style_cell(ws2.cell(r, 2, ecole), size=9)
    style_cell(ws2.cell(r, 3, prix), size=9, bold=True)
    style_cell(ws2.cell(r, 4, raison), size=9)
    lc = ws2.cell(r, 5)
    if lien.startswith("http"):
        lc.value = "Ouvrir ↗"
        lc.hyperlink = lien
        lc.font = Font(name=FONT, size=9, color=LINK_BLUE, underline="single")
        lc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        lc.border = border
    else:
        style_cell(lc, size=9)
        lc.value = lien
    ws2.row_dimensions[r].height = 90
    r += 1

# ============================================================
# ONGLET 3 — PARCOURS MÉTIER
# ============================================================
ws3 = wb.create_sheet("Parcours métier")
ws3.merge_cells("A1:B1")
t3 = ws3.cell(1, 1, "COMMENT ALLER VERS LES JOBS — Brand Partnerships Manager / Brand Strategist (mode & luxe)")
t3.font = Font(name=FONT, bold=True, size=12, color=WHITE)
t3.fill = PatternFill("solid", fgColor=NAVY)
t3.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 26
ws3.column_dimensions["A"].width = 32
ws3.column_dimensions["B"].width = 95

parcours = [
    ("Le job visé",
     "Brand Partnerships Manager (le plus exact) : gère les collabs entre marques, négocie les partenariats, choisit avec qui s'associer. Mix marketing + stratégie + négo. Variantes : Strategic Partnerships Manager (deals long terme), Collaborations Manager (collabs mode), Brand Strategist (positionnement/storytelling), Brand Manager/Director (gardien de l'identité)."),
    ("Compétences clés à bâtir",
     "1) Stratégie de marque (positionnement, identité, brand equity). 2) Négociation & business development (monter et closer des deals). 3) Marketing & communication (campagnes, storytelling). 4) Culture mode/luxe (connaître les maisons, les codes, les acteurs). 5) Gestion de projet & relationnel (gérer plusieurs partenaires). 6) Anglais courant indispensable."),
    ("Parcours type",
     "Bac+3/+5 marketing-communication-business OU mode → un master spécialisé branding/marque/mode (cf. onglet Comparatif) → premières expés : assistant·e brand / marketing / partnerships, agence de pub ou de relations marques, ou côté maison de mode. Puis montée vers Brand Partnerships / Brand Manager en 3-6 ans."),
    ("Premières expériences qui comptent",
     "Stages/jobs en : département marque ou partnerships d'une maison (LVMH, Kering, Richemont, marques DTC), agence de branding/pub, agence d'influence ou de partenariats, service presse/PR mode. Tout ce qui te met en contact avec la négociation marque×marque ou marque×personnalité."),
    ("Ce que regardent les recruteurs",
     "Un portfolio de projets concrets (même étudiants : une collab fictive bien construite vaut de l'or), un réseau dans le secteur (le master sert surtout à ça), et la preuve que tu sais faire le pont entre créa et business."),
    ("Stratégie recommandée (vu ta situation)",
     "Tu travailles déjà full-time → vise un master OFFICIEL et reconnu, en weekend Madrid ou online, qui te donne (a) la crédibilité du diplôme, (b) le réseau mode/luxe espagnol, (c) un projet/stage pour pivoter. Top combo : UPF-BSM (branding pur, réseau, 12k ferme) ou Universidad Europea moda online (officiel + mode + stage) ; ELLE MBA Lujo online (11 450 €) si tu veux le réseau luxe pur à Madrid."),
    ("Note de méthode",
     "Analyse issue d'une recherche multi-sources avec vérification adversariale des faits (17 faits confirmés à 3 voix sur 3) + 4 vérifications de prix ciblées. Les prix marqués « non publié » sont gardés derrière un formulaire par les écoles → voir l'onglet Notes & actions pour les obtenir."),
]
r = 2
for k, v in parcours:
    style_cell(ws3.cell(r, 1, k), fill=GREY, bold=True, size=10)
    style_cell(ws3.cell(r, 2, v), size=10)
    ws3.row_dimensions[r].height = max(60, 13 * (len(v) // 78 + 2))
    r += 1

# ============================================================
# ONGLET 4 — NOTES & ACTIONS
# ============================================================
ws4 = wb.create_sheet("Notes & actions")
heads4 = [("École / Master", 38), ("À demander / action", 50), ("Comment", 38)]
for j, (h, w) in enumerate(heads4, start=1):
    style_header(ws4.cell(1, j, h))
    ws4.column_dimensions[get_column_letter(j)].width = w
ws4.row_dimensions[1].height = 26

actions = [
    ("UPF-BSM — Brand Strategy (RECO n°1)", "Confirmer places REMOTE restantes (max 15/promo) + calendrier de candidature rentrée automne.", "Mail admissions BSM. Postuler TÔT — c'est le meilleur choix, à prioriser."),
    ("Universidad Europea — Máster Moda Online", "PRIX EXACT (non publié, fourchette 6-20k) + profiter de la promo -23% (échéance ~15 juin).", "Formulaire « solicitar información » → un conseiller donne le prix. Demander la brochure PDF."),
    ("UNIR — Gestión de Marca", "PRIX EXACT (non publié, est. 4-10k) + promo -30% (~2 juin). Confirmer rentrée nov. 2026.", "Formulaire UNIR ou appel ; demander total + mensualités."),
    ("ELLE Education — MBA & Diploma", "Dates de rentrée 2026/27 + bourses/early-bird + échéancier de paiement.", "Formulaire ELLE Education."),
    ("IED Madrid — Fashion Companies Mgmt", "Nb de soirs réels/sem + heure de fin (rapporté 18h30-22h30) → tenable avec ton job ? + frais d'inscription ?", "Mail/appel admissions IED. Point bloquant = le format soir semaine."),
    ("ESIC — MML (écarté : 20 400 €)", "SI vraiment intéressée : négocier les 30% d'aide (ramènerait ~14k). Sinon laisser (hors budget).", "Admissions ESIC — demander 'ayuda al estudio'."),
]
r = 2
for ecole, dem, comment in actions:
    style_cell(ws4.cell(r, 1, ecole), fill=GREY, bold=True, size=9)
    style_cell(ws4.cell(r, 2, dem), fill=YELLOW, size=9)
    style_cell(ws4.cell(r, 3, comment), size=9)
    ws4.row_dimensions[r].height = 55
    r += 1

# classements / hiérarchie note
r += 1
note = ws4.cell(r, 1, "Hiérarchie de prestige (pour info) : UPF-BSM & UCM (universités cotées QS) > Universidad Europea & UNIR (officiel, marque moyenne) > ELLE & IED (marque sectorielle mode/luxe, titres propres). ESIC & IE plus prestigieux mais hors budget 15k.")
ws4.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
note.font = Font(name=FONT, italic=True, size=9)
note.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws4.row_dimensions[r].height = 40

out = "/Users/bettydeclety/Downloads/Masters_Branding_Espagne_2026_v2.xlsx"
wb.save(out)
print("Saved:", out)
